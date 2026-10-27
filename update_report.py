import shutil
import sys
from pathlib import Path
from docx import Document

def replace_in_paragraph(paragraph, replacements):
    for run in paragraph.runs:
        for old, new in replacements.items():
            if old in run.text:
                run.text = run.text.replace(old, new)

def replace_in_table(table, replacements):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_paragraph(paragraph, replacements)

def process_doc(doc_path, replacements):
    doc = Document(doc_path)
    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        replace_in_table(table, replacements)
    doc.save(doc_path)

def main():
    if len(sys.argv) != 2:
        print("Usage: python update_report.py <docx_path>")
        sys.exit(1)
    docx_path = Path(sys.argv[1])
    replacements = {
        "OneShop": "UTE Tea",
        "ONE SHOP": "UTE TEA",
        "ONESHOP": "UTETE A",
        "OneShop - Báo cáo kiến trúc, sơ đồ hoạt động và giao diện": "UTE Tea - Báo cáo kiến trúc, sơ đồ hoạt động và giao diện",
        "OneShop - Báo cáo kiến trúc, sơ đồ hoạt động và giao diện": "UTE Tea - Báo cáo kiến trúc, sơ đồ hoạt động và giao diện",
        "OneShop - Báo cáo kiến trúc, sơ đồ hoạt động và giao diện": "UTE Tea - Báo cáo kiến trúc, sơ đồ hoạt động và giao diện"
    }
    process_doc(docx_path, replacements)
    print(f"Updated report: {docx_path}")

if __name__ == "__main__":
    main()
