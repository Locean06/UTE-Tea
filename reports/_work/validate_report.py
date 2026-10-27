from pathlib import Path
from docx import Document
from lxml import etree
from PIL import Image
import json,zipfile,hashlib,re
ROOT=Path(__file__).resolve().parents[2];W=ROOT/'reports/_work';P=ROOT/'reports/Bao_Cao_UTE_Tea.docx'
result={};hashes=json.loads((W/'source_hashes.json').read_text());result['java_files_checked']=len(hashes)
result['java_changes']=[p for p,h in hashes.items() if hashlib.sha256((ROOT/p).read_bytes()).hexdigest()!=h]
ex=json.loads((W/'excerpts.json').read_text());issues=[]
for e in ex:
 lines=(ROOT/e['source']).read_text().splitlines()[e['start']-1:e['end']]
 if hashlib.sha256('\n'.join(lines).encode()).hexdigest()!=e['sha256']:issues.append(e)
result['code_segments_verified']=len(ex);result['code_segment_issues']=issues
original=next(ROOT.glob('*.docx'));inv=json.loads((W/'document_inventory.json').read_text())
result['original_template_unchanged']=hashlib.sha256(original.read_bytes()).hexdigest()==inv['template_sha256']
d=Document(P);o=Document(original);result['chapters']=[p.text for p in d.paragraphs if p.style.name=='Heading 1'];result['sections']=len(d.sections)
result['page_setups_unchanged']=len(d.sections)==len(o.sections) and all(all(getattr(a,k)==getattr(b,k) for k in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin']) for a,b in zip(d.sections,o.sections))
result['paragraphs']=len(d.paragraphs);result['tables']=len(d.tables);result['inline_shapes']=len(d.inline_shapes)
with zipfile.ZipFile(P) as z,zipfile.ZipFile(original) as oz:
 result['zip_bad_member']=z.testzip();xml='\n'.join(z.read(n).decode('utf-8') for n in z.namelist() if n.endswith('.xml'))
 result['old_project_terms']=re.findall(r'OneShop|mỹ phẩm|24110248|24110269|24110368|NGÔ MINH KHÁNH|LÊ VĂN LÂN|ĐOÀN ANH TUẤN|WEPR330479_11',xml,re.I)
 oldmedia={hashlib.sha256(oz.read(n)).hexdigest() for n in oz.namelist() if n.startswith('word/media/')}
 result['unchanged_old_media']=[n for n in z.namelist() if n.startswith('word/media/') and hashlib.sha256(z.read(n)).hexdigest() in oldmedia]
 result['field_instructions']=len(re.findall('<w:instrText',xml));result['xml_parse_errors']=[]
 for n in z.namelist():
  if n.endswith('.xml'):
   try:etree.fromstring(z.read(n))
   except Exception:result['xml_parse_errors'].append(n)
 # Check that all final picture relationships resolve to actual package members.
 ns={'r':'http://schemas.openxmlformats.org/package/2006/relationships'}
 result['missing_image_targets']=[]
 for rel in etree.fromstring(z.read('word/_rels/document.xml.rels')).findall('r:Relationship',ns):
  if rel.get('Type','').endswith('/image'):
   target=rel.get('Target');target=target.lstrip('/') if target.startswith('/') else 'word/'+target
   if target not in z.namelist():result['missing_image_targets'].append(target)
result['diagram_png']=len(list((ROOT/'reports/So_Do').glob('*.png')));result['diagram_mermaid']=len(list((ROOT/'reports/So_Do').glob('*.mmd')));result['screenshots']=len(list((ROOT/'reports/Anh_Giao_Dien').glob('*.png')))
result['invalid_png']=[]
for p in list((ROOT/'reports/So_Do').glob('*.png'))+list((ROOT/'reports/Anh_Giao_Dien').glob('*.png')):
 try:
  with Image.open(p) as im:im.verify()
 except Exception:result['invalid_png'].append(str(p.relative_to(ROOT)))
# Cached caption counters are checked before any third-party field updater.
caption_groups={};bad=[]
for p in d.paragraphs:
 if p.style.name=='Caption':
  m=re.match(r'(Hình|Bảng|Đoạn mã) (\d+)\.(\d+)\.',p.text)
  if m:caption_groups.setdefault((m[1],m[2]),[]).append(int(m[3]))
for key,values in caption_groups.items():
 if values!=list(range(1,len(values)+1)):bad.append(str(key))
result['caption_sequence_errors']=bad
result['toc_fields']=len(re.findall(r'TOC \\',xml));result['docx_sha256']=hashlib.sha256(P.read_bytes()).hexdigest()
(W/'quality_checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False,indent=2))
assert not any(result[k] for k in ['java_changes','code_segment_issues','old_project_terms','unchanged_old_media','xml_parse_errors','missing_image_targets','invalid_png','caption_sequence_errors'])
assert result['original_template_unchanged'] and result['page_setups_unchanged'] and result['toc_fields']==4
